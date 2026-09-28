"""
Checks for the Review-II deck generator (offline).
Run:  python -m pytest tests/test_review2_deck.py -q
"""
import json
import math
import os
import re
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

from scripts.build_review2_deck import Ledger, build_slides, main

RESULTS = ["fall_engine_analysis.json", "ltmm_results.json", "ltmm_experiment2_results.json",
           "ltmm_daily_results.json", "fall_evaluation_results.json", "fall_evaluation_summary.json"]
needs_results = pytest.mark.skipif(not all(os.path.exists(os.path.join(ROOT, f)) for f in RESULTS),
                                   reason="results files not present")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("deck")
    assert main(["--out", str(out)]) == 0
    L = Ledger(ROOT)
    slides, qa = build_slides(L)
    return out, L, slides


@needs_results
def test_every_json_number_matches_its_source_key(built):
    _, L, _ = built
    keyed = [r for r in L.rows if r[4][0] == "key"]
    assert len(keyed) > 50
    for shown, *_rest, (kind, fname, keys, fmt) in keyed:
        d = json.load(open(os.path.join(ROOT, fname), encoding="utf-8"))
        for k in keys:
            d = d[k]
        assert not (isinstance(d, float) and math.isnan(d))
        assert shown == fmt.format(d), (fname, keys)


@needs_results
def test_agent_counts_match_the_records(built):
    _, L, _ = built
    recs = json.load(open(os.path.join(ROOT, "fall_evaluation_results.json"), encoding="utf-8"))
    for v in "ABCD":
        em = [r for r in recs if r["version"] == v and r["expected_emergency"]]
        expected = f"{sum(r['emergency_detected'] for r in em)}/{len(em)}"
        assert any(r[0] == expected and r[3] == f"{v} emergencies caught" for r in L.rows), v


@needs_results
def test_no_missing_values_and_no_hand_typed_result_numbers(built):
    out, L, _ = built
    ledger_nums = set()
    for r in L.rows:
        ledger_nums |= set(re.findall(r"\d+(?:\.\d+)?", r[0]))
    for name in ("SLIDES.md", "QA_PREP.md"):
        text = open(os.path.join(out, name), encoding="utf-8").read()
        assert "{{MISSING" not in text, name
        text = re.sub(r"!\[.*?\]\(.*?\)", "", text)                       # image paths
        text = re.sub(r"(?m)^- \[\d+\].*$", "", text)                     # IEEE reference entries
        text = re.sub(r"```.*?```", "", text, flags=re.S)                 # verbatim code excerpt
        text = re.sub(r"(?m)^\| \d+ \| .*$", "", text)                    # literature-table rows (row no. + author, year)
        text = re.sub(r"(?m)^\| \d\d–\d\d Oct 2026 .*$|^\| 28 Oct 2026 .*$", "", text)   # proposed timeline
        text = re.sub(r"\b(19|20)\d\d\b|COVID-19|30\.09\.2026|\(75%\)|95% CI|\(429\)|\[\d+\]", "", text)
        text = re.sub(r"(Rows|Papers|rows) \d+–\d+|Papers 1–10|Slide \d+ —|slide \d+|\(\d/\d\)", "", text)
        stray = sorted({m for m in re.findall(r"(?<![\w.])\d+(?:\.\d+)?", text) if m not in ledger_nums})
        assert not stray, f"{name}: numbers not traceable to the ledger: {stray}"


@needs_results
def test_pptx_keeps_template_placeholders_and_has_notes(built):
    from pptx import Presentation
    out, _, slides = built
    prs = Presentation(os.path.join(out, "Review-II_MedXAI.pptx"))
    assert len(prs.slides) == len(slides)
    texts = ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert "Name — Reg. No." in texts[0] and "Dr. Guide Name, School" in texts[0] and "Signature:" in texts[0]
    for i, s in enumerate(prs.slides, 1):
        assert s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip(), i
        assert "Add your results here" not in s.notes_slide.notes_text_frame.text
        nums = [sh.text_frame.text for sh in s.shapes if sh.name == "Text 0"]
        assert nums == [str(i)], i                                      # page numbers are sequential
    alltext = "\n".join(texts)
    assert "A. Author, B. Author" not in alltext and "Project Title" not in alltext
